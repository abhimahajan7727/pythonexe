import customtkinter as ctk
import tkinter as tk
from tkinter import messagebox
from datetime import datetime, timedelta
import pyperclip, os, pandas as pd

# --- Data Helpers (unchanged) ---
default_excel   = "D:/Programs/Code/zoom/Monthly_accexcel.xlsx"
regular_excel   = "D:/Programs/Code/zoom/Regular_accexcel.xlsx"
seller_file     = "sellers.txt"

def load_sellers():
    if not os.path.exists(seller_file):
        return []
    with open(seller_file, 'r') as f:
        return sorted({line.strip() for line in f if line.strip()})

def save_seller(name):
    with open(seller_file, 'a') as f:
        f.write(name + '\n')

# --- Excel Helpers (unchanged) ---
def ensure_excel(path):
    if not os.path.exists(path):
        pd.DataFrame(columns=[
            "Zoom Mail","Zoom Password",
            "Expiry Date","Expiry Time",
            "Plan","Seller","Renew Count"
        ]).to_excel(path, index=False)

def update_excel(mail, pw, exp_date, exp_time, plan, seller, renew, path):
    ensure_excel(path)
    df = pd.read_excel(path)
    if mail in df["Zoom Mail"].values:
        idx = df[df["Zoom Mail"] == mail].index[0]
        df.loc[idx, ["Zoom Password","Expiry Date","Expiry Time","Plan","Seller"]] = [
            pw, exp_date, exp_time, plan, seller
        ]
        if renew:
            df.at[idx, "Renew Count"] += 1
    else:
        df.loc[len(df)] = [
            mail, pw, exp_date, exp_time, plan, seller,
            1 if renew else 0
        ]
    df.to_excel(path, index=False)

    # --- Format the Excel file after saving ---
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    ws = wb.active

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val_len = len(str(cell.value))
                if val_len > max_length:
                    max_length = val_len
            except:
                pass
        ws.column_dimensions[col_letter].width = max_length + 2

    wb.save(path)


# --- Enhanced AutoComplete Entry ---
class AutoCompleteEntry(ctk.CTkFrame):
    def __init__(self, master, values, **kwargs):
        super().__init__(master, **kwargs)
        self.values = values
        self.var = ctk.StringVar()
        self.entry = ctk.CTkEntry(self, textvariable=self.var, corner_radius=10)
        self.entry.pack(side="left", fill="x", expand=True)

        self.toggle_btn = ctk.CTkButton(
        self,
        text="▾",
        width=30,
        corner_radius=8,
        fg_color="transparent",        # still transparent fill
        text_color="#333333",          # dark arrow so it shows on light bg
        border_width=1,                # draw a thin border
        border_color="#888888",        # light-gray border
        hover_color="#dddddd",         # hover highlight
        command=self.toggle_suggestions)
        self.toggle_btn.pack(side="right", padx=(5,0))

        self.dropdown = tk.Toplevel(self)
        self.dropdown.withdraw()
        self.dropdown.overrideredirect(True)
        self.dropdown.configure(bg="white")

        self.listbox = tk.Listbox(
            self.dropdown, height=0, font=("Segoe UI", 10)
        )
        self.scrollbar = tk.Scrollbar(
            self.dropdown, command=self.listbox.yview
        )
        self.listbox.configure(yscrollcommand=self.scrollbar.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        self.scrollbar.pack(side="right", fill="y")

        self.listbox.bind("<<ListboxSelect>>", self.fill_from_list)
        self.entry.bind("<KeyRelease>", self.on_keyrelease)
        self.listbox.bind("<FocusOut>", self.hide_suggestions)
        self.entry.bind("<Down>", lambda e: self._focus_list())
        self.entry.bind("<Tab>", self.select_first_item)

    def on_keyrelease(self, event):
        text = self.var.get().strip().lower()
        matches = [v for v in self.values if v.lower().startswith(text)] if text else []
        self.listbox.delete(0, "end")
        if matches:
            for m in matches:
                self.listbox.insert("end", m)
            self.show_suggestions()
        else:
            self.hide_suggestions()

    def toggle_suggestions(self):
        if self.dropdown.winfo_viewable():
            self.hide_suggestions()
        else:
            self.listbox.delete(0, "end")
            for v in self.values:
                self.listbox.insert("end", v)
            self.show_suggestions()

    def show_suggestions(self):
        count = self.listbox.size()
        if count == 0:
            return
        x = self.entry.winfo_rootx()
        y = self.entry.winfo_rooty() + self.entry.winfo_height()
        width = self.entry.winfo_width() + self.toggle_btn.winfo_width() + 5
        height = min(5, count) * 24
        self.dropdown.geometry(f"{width}x{height}+{x}+{y}")
        self.dropdown.deiconify()

    def hide_suggestions(self, *_):
        self.dropdown.withdraw()

    def fill_from_list(self, event):
        if not self.listbox.curselection():
            return
        sel = self.listbox.get(self.listbox.curselection()[0])
        self.var.set(sel)
        self.hide_suggestions()

    def select_first_item(self, event):
        if self.listbox.size():
            self.var.set(self.listbox.get(0))
            self.hide_suggestions()
        return "break"

    def _focus_list(self):
        if self.listbox.size():
            self.listbox.focus_set()
            self.listbox.selection_set(0)

    def get(self):
        return self.var.get().strip()

    def update_values(self, values):
        self.values = values

# --- Formatting Logic ---
def format_accounts(text, plan, months, seller, renew):
    now = datetime.now()
    excel_path = regular_excel if plan == "Regular" else default_excel
    out_lines = []

    for block in filter(bool, text.split("\n\n")):
        mail, pw = "", "Zoom@123"
        for line in block.splitlines():
            low = line.lower().strip()
            if low.startswith("mail:"):
                mail = line.split(":",1)[1].strip()
            if low.startswith("password:"):
                val = line.split(":",1)[1].strip()
                pw  = val or pw
        if not mail:
            continue

        if plan in ("Regular", "14 days"):
            days = 14
        elif plan in ("1 Month", "28 days"):
            days = 28
        else:
            days = 28 

        exp_dt = now + timedelta(days=days)
        exp_clip = exp_dt.strftime("%b %d, %Y")
        time_clip = exp_dt.strftime("%I:%M %p")
        plan_name = plan if plan != "Custom" else f"{months} Months"

        update_excel(mail, pw, exp_clip, time_clip, plan_name, seller, renew, excel_path)

        out_lines += [
            "Zoom - Zoom website",
            f"Zoom Mail: {mail}",
            f"Zoom password: {pw}",
            "",
            f"Next Renew Date: {exp_clip}",
            "co-host: On",
            "YouTube Stream: On",
            ""
        ]

    result = "\n".join(out_lines)
    pyperclip.copy(result)
    return result

# --- GUI ---
ctk.set_appearance_mode("System")
ctk.set_default_color_theme("dark-blue")

class ZoomFormatter(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Zoom Account Formatter")
        self.minsize(850, 620)
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)

        # Header
        header = ctk.CTkFrame(self, fg_color="#ececec", corner_radius=0)
        header.grid(row=0, column=0, sticky="ew")
        header.grid_columnconfigure((0,1,2), weight=1)
        ctk.CTkLabel(
            header, text="Zoom Account Formatter",
            font=("Segoe UI", 24, "bold"), text_color="#333333"
        ).grid(row=0, column=0, padx=20, pady=15, sticky="w")
        self.last_label = ctk.CTkLabel(
            header, text="Last updated: never",
            font=("Segoe UI", 12), text_color="#555555"
        )
        self.last_label.grid(row=0, column=1)
        ctk.CTkButton(
            header, text="⚙️", width=30, height=30,
            corner_radius=15, fg_color="transparent",
            hover_color="#dddddd", command=self.open_settings
        ).grid(row=0, column=2, padx=20)

        # Body
        body = ctk.CTkFrame(self, fg_color="transparent")
        body.grid(row=1, column=0, padx=20, pady=10, sticky="nsew")
        body.grid_columnconfigure(0, weight=1)
        body.grid_rowconfigure(1, weight=1)

        # Credentials Section
        creds = ctk.CTkFrame(body, fg_color="#ffffff", corner_radius=10)
        creds.grid(row=0, column=0, sticky="ew", pady=(0,10))
        creds.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            creds, text="🔐 Account Credentials",
            font=("Segoe UI", 18, "bold")
        ).grid(row=0, column=0, padx=15, pady=(10,5), sticky="w")
        self.textbox = ctk.CTkTextbox(
            creds, corner_radius=10, font=("Segoe UI", 12)
        )
        self.textbox.grid(row=1, column=0, padx=15, pady=(0,15), sticky="nsew")
        self.textbox.insert(
            "0.0",
            "Mail:\nPassword: Zoom@123\nRecovery Mail: codestream.zoom@manyme.com"
        )

        # Subscription Section
        sub = ctk.CTkFrame(body, fg_color="#ffffff", corner_radius=10)
        sub.grid(row=1, column=0, sticky="nsew")
        sub.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            sub, text="📝 Subscription Details",
            font=("Segoe UI", 18, "bold")
        ).grid(row=0, column=0, padx=15, pady=(10,5), sticky="w")
        form = ctk.CTkFrame(sub, fg_color="transparent")
        form.grid(row=1, column=0, padx=15, pady=10, sticky="ew")
        form.grid_columnconfigure(1, weight=1)

        # Seller Row
        ctk.CTkLabel(
            form, text="👤 Seller", font=("Segoe UI", 12, "bold")
        ).grid(row=0, column=0, sticky="w")
        self.seller_box = AutoCompleteEntry(form, values=load_sellers())
        self.seller_box.grid(row=0, column=1, sticky="ew", padx=5)
        ctk.CTkButton(
            form, text="Add", width=80,
            fg_color="#0288d1", hover_color="#03a9f4",
            command=self.add_seller
        ).grid(row=0, column=2, padx=5)

        # Plan & Custom Months
        ctk.CTkLabel(
            form, text="📦 Plan Type", font=("Segoe UI", 12, "bold")
        ).grid(row=1, column=0, pady=8, sticky="w")
        self.plan_var = ctk.StringVar(value="1 Month")
        for idx, p in enumerate(["1 Month","Regular","Custom"]):
            ctk.CTkRadioButton(
                form, text=p, variable=self.plan_var, value=p
            ).grid(row=1, column=1+idx, padx=5, sticky="w")
        ctk.CTkLabel(
            form, text="📅 Custom Months", font=("Segoe UI", 12, "bold")
        ).grid(row=2, column=0, pady=8, sticky="w")
        self.months_var = ctk.StringVar(value="1")
        ctk.CTkComboBox(
            form, values=[str(i) for i in range(1,13)],
            variable=self.months_var, width=80
        ).grid(row=2, column=1, sticky="w", padx=5)
        self.renew_var = ctk.BooleanVar()
        ctk.CTkCheckBox(
            form, text="Renew Existing Accounts",
            variable=self.renew_var, font=("Segoe UI", 12)
        ).grid(row=3, column=0, columnspan=2, pady=12, sticky="w")

        # Footer - Process Button
        footer = ctk.CTkFrame(self, fg_color="transparent")
        footer.grid(row=2, column=0, pady=(10,20))
        process_btn = ctk.CTkButton(
            footer, text="✅ Process & Copy", width=220, height=50,
            fg_color="#0288d1", hover_color="#03a9f4",
            font=("Segoe UI", 16, "bold"),
            command=self.process_and_copy
        )
        process_btn.pack()

    def add_seller(self):
        name = self.seller_box.get()
        if not name:
            messagebox.showwarning("Input Required", "Enter a seller name.")
            return
        current = load_sellers()
        if name not in current:
            save_seller(name)
            self.seller_box.update_values(sorted(current))
            messagebox.showinfo("Seller Added", f"'{name}' added.")

    def process_and_copy(self):
        text = self.textbox.get("0.0", "end").strip()
        seller = self.seller_box.get()
        if not text:
            messagebox.showerror("Missing Input", "Please enter account data.")
            return
        if not seller:
            messagebox.showerror("Missing Seller", "Please select or enter a seller.")
            return

        spinner = ctk.CTkProgressBar(self, mode="indeterminate", width=200)
        spinner.place(relx=0.5, rely=0.05, anchor="n")
        spinner.start()

        format_accounts(
            text, self.plan_var.get(), self.months_var.get(),
            seller, self.renew_var.get()
        )
        spinner.stop()
        spinner.destroy()
        self.last_label.configure(
            text=f"Last updated: {datetime.now().strftime('%H:%M:%S')}"
        )
        messagebox.showinfo("Success", "Accounts processed and copied!")

    def open_settings(self):
        dlg = ctk.CTkToplevel(self)
        dlg.title("Settings")
        dlg.minsize(300, 200)
        mode_var = ctk.StringVar(value=ctk.get_appearance_mode())
        ctk.CTkLabel(dlg, text="Appearance Mode", font=("Segoe UI", 14)).pack(pady=10)
        mode_menu = ctk.CTkOptionMenu(
            dlg, values=["Light", "Dark", "System"], variable=mode_var
        )
        mode_menu.pack(pady=5)
        ctk.CTkButton(
            dlg, text="Apply", command=lambda: (ctk.set_appearance_mode(mode_var.get()), dlg.destroy())
        ).pack(pady=20)

if __name__ == "__main__":
    ZoomFormatter().mainloop()